"""Planilha de saída `respostas_<simulado>.xlsx` e `relatorio.json` (PRD 7.4, com a matrícula na coluna 1).

A ordem das abas importa: a tela de importação do Academy lê **só a primeira**, por posição: coluna 1 =
Matrícula/RA, colunas 2 a N+1 = questões 1 a N, na ordem, e pula a linha de cabeçalho. Por isso a aba
`Respostas` tem exatamente `1 + N` colunas e nada além disso.
"""
from __future__ import annotations

import json
from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from . import __version__
from .models import sim8
from .processar import (
    ALUNO_DUPLICADO,
    IGNORADA,
    NAO_IDENTIFICADOS,
    RESPOSTAS,
    PaginaLida,
    Resultado,
)

# 1ª coluna: matrícula/RA ou, no plano B (aluno sem matrícula no cadastro), o e-mail. O cabeçalho é só para quem lê;
# a importação lê a coluna por posição.
COLUNA_ID = "matricula_ra_ou_email"
AMARELO = PatternFill("solid", fgColor="FFF2CC")
NEGRITO = Font(bold=True)


def _cabecalho(n: int) -> list[str]:
    return [COLUNA_ID] + [f"Questão {q}" for q in range(1, n + 1)]


def _linha(pg: PaginaLida, n: int) -> list[str | None]:
    return [pg.identificacao or ""] + [pg.respostas.get(q) for q in range(1, n + 1)]


def _escreve_respostas(ws, paginas: list[PaginaLida], n: int, com_pagina: bool = False) -> None:
    deslocamento = 2 if com_pagina else 0
    cab = (["pagina", "motivo"] if com_pagina else []) + _cabecalho(n)
    ws.append(cab)
    for c in ws[1]:
        c.font = NEGRITO
    for pg in paginas:
        linha = ([pg.pagina, pg.motivo] if com_pagina else []) + _linha(pg, n)
        ws.append(linha)
        r = ws.max_row
        ws.cell(r, deslocamento + 1).number_format = "@"  # matrícula é texto: preserva zeros à esquerda
        for q, tipo in pg.duvidas.items():
            cel = ws.cell(r, deslocamento + 1 + q)
            cel.fill = AMARELO
            cel.comment = Comment(f"{'múltipla' if tipo == 'multipla' else 'ambígua'} (página {pg.pagina})", "cartao-omr")
    ws.freeze_panes = ws.cell(2, deslocamento + 2)
    ws.column_dimensions["A"].width = 22 if not com_pagina else 9
    for col in range(deslocamento + 2, deslocamento + n + 2):
        ws.column_dimensions[get_column_letter(col)].width = 11


def resumo(res: Resultado) -> dict:
    cont = Counter(pg.status for pg in res.paginas)
    destinos = Counter(pg.destino for pg in res.paginas)
    celulas = Counter()
    for pg in res.paginas:
        if pg.destino == IGNORADA:
            continue
        marcadas = sum(1 for v in pg.respostas.values() if v)
        duvidas = Counter(pg.duvidas.values())
        celulas["marcadas"] += marcadas
        celulas["multiplas"] += duvidas.get("multipla", 0)
        celulas["ambiguas"] += duvidas.get("ambigua", 0)
        celulas["em_branco"] += len(pg.respostas) - marcadas - sum(duvidas.values())
    return {
        "simulado": res.simulado_id,
        "sim8": sim8(res.simulado_id),
        "arquivo": res.arquivo,
        "executado_em": datetime.now().isoformat(timespec="seconds"),
        "cartao_omr": __version__,
        "total_questoes": res.total_questoes,
        "limiares": {"T_ALTO": res.cfg.t_alto, "T_BAIXO": res.cfg.t_baixo, "GAP_MIN": res.cfg.gap_min,
                     "T_INK": res.cfg.t_ink},
        "paginas": len(res.paginas),
        "em_branco": cont.get("pagina_em_branco", 0),
        "ilegiveis": cont.get("pagina_ilegivel", 0),
        "simulado_divergente": cont.get("simulado_divergente", 0),
        "identificadas": destinos.get(RESPOSTAS, 0),
        "identificadas_por_email": sum(1 for pg in res.paginas if pg.destino == RESPOSTAS and pg.via == "email"),
        "nao_identificadas": destinos.get(NAO_IDENTIFICADOS, 0),
        "duplicadas": sum(1 for pg in res.paginas if pg.motivo == ALUNO_DUPLICADO),
        "questoes": dict(celulas),
        "pendencias": len(res.pendencias),
        "ausentes_da_lista": len(res.ausentes),
    }


def escrever_planilha(res: Resultado, dest: Path) -> Path:
    n = res.total_questoes
    wb = Workbook()
    ws = wb.active
    ws.title = "Respostas"
    ok = sorted((pg for pg in res.paginas if pg.destino == RESPOSTAS), key=lambda p: (p.identificacao or "", p.pagina))
    _escreve_respostas(ws, ok, n)

    wp = wb.create_sheet("Pendências")
    wp.append(["pagina", "tipo", "questao", "detalhe", "recorte"])
    for c in wp[1]:
        c.font = NEGRITO
    for p in sorted(res.pendencias, key=lambda p: (p.pagina, p.questao or 0)):
        wp.append([p.pagina, p.tipo, p.questao, p.detalhe, p.recorte or ""])
    for col, w in zip("ABCDE", (8, 24, 9, 60, 40)):
        wp.column_dimensions[col].width = w

    wn = wb.create_sheet("Não identificados")
    _escreve_respostas(wn, [pg for pg in res.paginas if pg.destino == NAO_IDENTIFICADOS], n, com_pagina=True)

    wr = wb.create_sheet("Resumo")
    for chave, valor in resumo(res).items():
        if isinstance(valor, dict):
            for k, v in valor.items():
                wr.append([f"{chave}.{k}", v])
        else:
            wr.append([chave, valor])
    wr.column_dimensions["A"].width = 28
    wr.column_dimensions["B"].width = 40
    if res.ausentes:
        wr.append([])
        wr.append(["da lista da emissão, sem cartão devolvido"])
        for m in res.ausentes:
            wr.append([m])

    dest.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dest)
    return dest


def escrever_relatorio(res: Resultado, dest: Path) -> Path:
    dados = {
        "resumo": resumo(res),
        "pendencias": [vars(p) for p in res.pendencias],
        "ausentes_da_lista": res.ausentes,
        "paginas": [
            {"pagina": pg.pagina, "destino": pg.destino, "status": pg.status, "layout": pg.layout,
             "identificacao": pg.identificacao, "via": pg.via, "motivo": pg.motivo}
            for pg in res.paginas
        ],
    }
    dest.write_text(json.dumps(dados, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return dest
